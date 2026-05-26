from __future__ import annotations

import json
import os
import re
import site
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

from analytics import build_grade_lookup, is_absent_grade, is_passing_grade


def _bootstrap_optional_packages() -> None:
    candidates: list[Path] = []
    env_path = os.environ.get("RESULT_ANALYSER_PYTHON_PACKAGES")
    if env_path:
        candidates.append(Path(env_path))

    for candidate in candidates:
        if not candidate.exists():
            continue
        site.addsitedir(str(candidate))
        site_packages = candidate / "Lib" / "site-packages"
        if site_packages.exists():
            site.addsitedir(str(site_packages))


try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    _bootstrap_optional_packages()
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
HEADER_IMAGE_PATH = BASE_DIR / "static" / "assets" / "aset-header.png"
WORD_BUILDER_PATH = BASE_DIR / "scripts" / "build_word_report.mjs"
WORD_BUILDER_PY_PATH = BASE_DIR / "scripts" / "build_word_report.py"

ACCENT_BLUE = "1157A7"
ACCENT_RED = "D34040"
ACCENT_GREEN = "2F6B5F"
LIGHT_FILL = "F4F7FB"
THIN_SIDE = Side(style="thin", color="D1D9E0")
TABLE_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

def slugify_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value.strip()).strip("-")
    return cleaned.lower() or "result-analysis"


def build_department_rankings(
    students: list[dict[str, Any]], rank_limit: int = 10
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for student in students:
        grouped[student["department_name"]].append(student)

    ranking_rows: list[dict[str, Any]] = []
    for department_name in sorted(grouped):
        ranked_students = sorted(
            grouped[department_name],
            key=lambda row: (-row["gpa"], row["register_no"]),
        )[:rank_limit]
        for index, student in enumerate(ranked_students, start=1):
            ranking_rows.append(
                {
                    "department_name": department_name,
                    "rank": index,
                    "register_no": student["register_no"],
                    "gpa": float(student["gpa"]),
                    "passed_count": int(student["passed_count"]),
                    "failed_count": int(student["failed_count"]),
                    "subject_count": int(student["subject_count"]),
                }
            )
    return ranking_rows


def build_failed_student_rows(
    students: list[dict[str, Any]], grade_config: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    grade_lookup = build_grade_lookup(grade_config)
    failed_rows: list[dict[str, Any]] = []

    for student in sorted(
        students,
        key=lambda row: (row["department_name"], row["register_no"]),
    ):
        failed_subjects = [
            f"{grade['course_code']} ({grade['grade']})"
            for grade in student["grades"]
            if not is_passing_grade(grade["grade"], grade_lookup)
            and not is_absent_grade(grade["grade"])
        ]
        if not failed_subjects:
            continue

        failed_rows.append(
            {
                "department_name": student["department_name"],
                "register_no": student["register_no"],
                "gpa": float(student["gpa"]),
                "failed_subjects": failed_subjects,
                "failed_subjects_text": ", ".join(failed_subjects),
                "failed_subject_count": len(failed_subjects),
            }
        )

    return failed_rows


def build_threshold_rows(
    students: list[dict[str, Any]], threshold: float
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    matching_students = sorted(
        [student for student in students if float(student["gpa"]) > threshold],
        key=lambda row: (-row["gpa"], row["register_no"]),
    )

    counts_by_department: list[dict[str, Any]] = []
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for student in matching_students:
        grouped[student["department_name"]].append(student)

    for department_name in sorted(grouped):
        counts_by_department.append(
            {
                "department_name": department_name,
                "count": len(grouped[department_name]),
            }
        )

    return matching_students, counts_by_department


def build_report_context(
    analytics: dict[str, Any],
    grade_config: list[dict[str, Any]],
    threshold: float,
) -> dict[str, Any]:
    students = analytics["students"]
    threshold_students, threshold_counts = build_threshold_rows(students, threshold)

    return {
        "batch": analytics["batch"],
        "summary": analytics["summary"],
        "students": students,
        "subjects": analytics["subjects"],
        "top_students": analytics["top_students"],
        "bottom_students": analytics["bottom_students"],
        "department_rankings": build_department_rankings(students, rank_limit=10),
        "failed_students": build_failed_student_rows(students, grade_config),
        "threshold": float(threshold),
        "threshold_students": threshold_students,
        "threshold_counts": threshold_counts,
        "grade_config": grade_config,
    }


def excel_filename(batch_name: str) -> str:
    return f"{slugify_filename(batch_name)}-analysis.xlsx"


def word_filename(batch_name: str) -> str:
    return f"{slugify_filename(batch_name)}-analysis.docx"


def resolve_node_modules_dir() -> str | None:
    return os.environ.get("RESULT_ANALYSER_NODE_MODULES")


def resolve_bundled_python_binary() -> str | None:
    return os.environ.get("RESULT_ANALYSER_PYTHON_BINARY")


def normalise_word_media_extensions(docx_bytes: bytes) -> bytes:
    image_extension = HEADER_IMAGE_PATH.suffix.lower() or ".png"
    with zipfile.ZipFile(BytesIO(docx_bytes), "r") as source_zip:
        files = {name: source_zip.read(name) for name in source_zip.namelist()}

    renamed_targets: dict[str, str] = {}
    for name in list(files):
        if not name.startswith("word/media/") or not name.endswith(".undefined"):
            continue
        new_name = name.removesuffix(".undefined") + image_extension
        files[new_name] = files.pop(name)
        renamed_targets[name.split("/", 1)[1]] = new_name.split("/", 1)[1]

    if not renamed_targets:
        return docx_bytes

    for name, content in list(files.items()):
        if not name.endswith(".rels"):
            continue
        updated_text = content.decode("utf-8")
        for old_target, new_target in renamed_targets.items():
            updated_text = updated_text.replace(old_target, new_target)
        files[name] = updated_text.encode("utf-8")

    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
        for name, content in files.items():
            target_zip.writestr(name, content)

    return output.getvalue()


def generate_excel_report(
    analytics: dict[str, Any],
    grade_config: list[dict[str, Any]],
    threshold: float,
) -> tuple[bytes, str]:
    context = build_report_context(analytics, grade_config, threshold)

    workbook = Workbook()
    overview_sheet = workbook.active
    overview_sheet.title = "Overview"
    workbook.create_sheet("Subject Pass %")
    workbook.create_sheet("Department Top 10")
    workbook.create_sheet("Failed Students")
    workbook.create_sheet("Above Threshold")
    workbook.create_sheet("Student GPA")

    build_overview_sheet(workbook["Overview"], context)
    build_subject_sheet(workbook["Subject Pass %"], context)
    build_department_sheet(workbook["Department Top 10"], context)
    build_failed_students_sheet(workbook["Failed Students"], context)
    build_threshold_sheet(workbook["Above Threshold"], context)
    build_student_sheet(workbook["Student GPA"], context)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), excel_filename(context["batch"]["name"])


def register_excel_filename(batch_name: str) -> str:
    return f"{slugify_filename(batch_name)}-gpa-register.xlsx"


def generate_student_register_excel(
    analytics: dict[str, Any],
    department: str = "",
    search: str = "",
) -> tuple[bytes, str]:
    students = analytics["students"]
    if department and department != "All Departments":
        students = [
            student
            for student in students
            if student["department_name"] == department
        ]

    search_key = search.strip().upper()
    if search_key:
        students = [
            student
            for student in students
            if search_key in student["register_no"].upper()
        ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GPA Register"
    start_row = prepare_excel_sheet(
        sheet,
        "GPA Register",
        analytics["batch"],
        "Student rows match the current dashboard register view.",
    )
    rows = [
        [
            student["register_no"],
            student["department_name"],
            student["gpa"],
            student["passed_count"],
            student["failed_count"],
            student["absent_count"],
            student["subject_count"],
            ", ".join(
                f"{grade['course_code']} ({grade['grade']})"
                for grade in student["grades"]
            ),
        ]
        for student in students
    ]
    write_table(
        sheet,
        start_row,
        [
            "Register No",
            "Department",
            "GPA",
            "Passed",
            "Failed",
            "Absent",
            "Total Subjects",
            "Grades",
        ],
        rows or [["No students match the current view", "", "", "", "", "", "", ""]],
    )
    autosize_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), register_excel_filename(analytics["batch"]["name"])


def build_overview_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "KTU Result Analysis Overview",
        context["batch"],
        f"GPA threshold report uses values greater than {context['threshold']:.2f}.",
    )

    summary = context["summary"]
    metrics = [
        ("Batch Name", context["batch"]["name"]),
        ("Exam Name", context["batch"]["exam_name"]),
        ("Institution", context["batch"]["institution"] or "Institution not detected"),
        ("Total Students", summary["total_students"]),
        ("Total Departments", summary["total_departments"]),
        ("Total Courses", summary["total_courses"]),
        ("Average GPA", summary["average_gpa"]),
        ("Highest GPA", summary["highest_gpa"]),
        ("Lowest GPA", summary["lowest_gpa"]),
        ("Students with GPA > threshold", len(context["threshold_students"])),
        (
            "Unknown Grades",
            ", ".join(summary["unknown_grades"]) if summary["unknown_grades"] else "None",
        ),
    ]
    write_key_value_block(sheet, start_row, metrics)

    top_row = start_row + len(metrics) + 3
    sheet.cell(row=top_row, column=1, value="Top Performers")
    style_section_title(sheet.cell(row=top_row, column=1))
    write_table(
        sheet,
        top_row + 1,
        ["Rank", "Register No", "Department", "GPA"],
        [
            [index, student["register_no"], student["department_name"], student["gpa"]]
            for index, student in enumerate(context["top_students"], start=1)
        ],
    )

    bottom_row = top_row + max(len(context["top_students"]), 1) + 5
    sheet.cell(row=bottom_row, column=1, value="Lowest GPA")
    style_section_title(sheet.cell(row=bottom_row, column=1))
    write_table(
        sheet,
        bottom_row + 1,
        ["Rank", "Register No", "Department", "GPA"],
        [
            [index, student["register_no"], student["department_name"], student["gpa"]]
            for index, student in enumerate(context["bottom_students"], start=1)
        ],
    )

    autosize_sheet(sheet)


def build_subject_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "Subject-wise Pass Percentage",
        context["batch"],
        "This sheet lists appeared, passed, failed, absent, and pass percentage for each subject.",
    )
    rows = [
        [
            subject["department_name"],
            subject["course_code"],
            subject["course_name"],
            subject["credits"],
            subject["appeared_count"],
            subject["passed_count"],
            subject["failed_count"],
            subject["absent_count"],
            subject["pass_percentage"],
        ]
        for subject in context["subjects"]
    ]
    write_table(
        sheet,
        start_row,
        [
            "Department",
            "Course Code",
            "Course Name",
            "Credits",
            "Appeared",
            "Passed",
            "Failed",
            "Absent",
            "Pass %",
        ],
        rows,
    )
    autosize_sheet(sheet)


def build_department_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "Top 10 Rank Holders in Each Department",
        context["batch"],
        "Ranks are computed department-wise using GPA in descending order.",
    )
    rows = [
        [
            row["department_name"],
            row["rank"],
            row["register_no"],
            row["gpa"],
            row["passed_count"],
            row["failed_count"],
            row["subject_count"],
        ]
        for row in context["department_rankings"]
    ]
    write_table(
        sheet,
        start_row,
        [
            "Department",
            "Rank",
            "Register No",
            "GPA",
            "Passed Subjects",
            "Failed Subjects",
            "Total Subjects",
        ],
        rows,
    )
    autosize_sheet(sheet)


def build_failed_students_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "Students with Failed Subjects",
        context["batch"],
        "Absent subjects are not included in the failed-subject list.",
    )
    rows = [
        [
            row["department_name"],
            row["register_no"],
            row["gpa"],
            row["failed_subject_count"],
            row["failed_subjects_text"],
        ]
        for row in context["failed_students"]
    ]
    write_table(
        sheet,
        start_row,
        [
            "Department",
            "Register No",
            "GPA",
            "Failed Subject Count",
            "Failed Subjects",
        ],
        rows or [["No failed students found", "", "", "", ""]],
    )
    autosize_sheet(sheet)


def build_threshold_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "Students Above GPA Threshold",
        context["batch"],
        f"This analysis counts students with GPA greater than {context['threshold']:.2f}.",
    )
    sheet.cell(row=start_row, column=1, value="Department-wise Count")
    style_section_title(sheet.cell(row=start_row, column=1))
    write_table(
        sheet,
        start_row + 1,
        ["Department", "Count"],
        [
            [row["department_name"], row["count"]]
            for row in context["threshold_counts"]
        ]
        or [["No students exceed the threshold", 0]],
    )

    detail_row = start_row + max(len(context["threshold_counts"]), 1) + 5
    sheet.cell(row=detail_row, column=1, value="Students Above Threshold")
    style_section_title(sheet.cell(row=detail_row, column=1))
    write_table(
        sheet,
        detail_row + 1,
        ["Register No", "Department", "GPA", "Passed", "Failed"],
        [
            [
                student["register_no"],
                student["department_name"],
                student["gpa"],
                student["passed_count"],
                student["failed_count"],
            ]
            for student in context["threshold_students"]
        ]
        or [["No students exceed the threshold", "", "", "", ""]],
    )
    autosize_sheet(sheet)


def build_student_sheet(sheet, context: dict[str, Any]) -> None:
    start_row = prepare_excel_sheet(
        sheet,
        "Complete Student GPA Register",
        context["batch"],
        "All subjects and grades are listed for every student in the batch.",
    )
    rows = [
        [
            student["register_no"],
            student["department_name"],
            student["gpa"],
            student["passed_count"],
            student["failed_count"],
            student["absent_count"],
            student["subject_count"],
            ", ".join(
                f"{grade['course_code']} ({grade['grade']})" for grade in student["grades"]
            ),
        ]
        for student in context["students"]
    ]
    write_table(
        sheet,
        start_row,
        [
            "Register No",
            "Department",
            "GPA",
            "Passed",
            "Failed",
            "Absent",
            "Total Subjects",
            "Grades",
        ],
        rows,
    )
    autosize_sheet(sheet)


def prepare_excel_sheet(
    sheet,
    title: str,
    batch: dict[str, Any],
    subtitle: str,
) -> int:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    for row_index in range(1, 5):
        sheet.row_dimensions[row_index].height = 18

    add_excel_header_image(sheet)

    sheet.merge_cells("A5:F5")
    title_cell = sheet["A5"]
    title_cell.value = title
    title_cell.font = Font(bold=True, size=18, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=ACCENT_BLUE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[5].height = 28

    sheet.merge_cells("A6:F6")
    subtitle_cell = sheet["A6"]
    subtitle_cell.value = (
        f"{batch['name']} | {batch['exam_name']} | {subtitle}"
    )
    subtitle_cell.font = Font(size=10, color="3F4D5D")
    subtitle_cell.fill = PatternFill("solid", fgColor=LIGHT_FILL)
    subtitle_cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[6].height = 32
    return 8


def add_excel_header_image(sheet) -> None:
    if not HEADER_IMAGE_PATH.exists():
        return

    try:
        image = OpenpyxlImage(str(HEADER_IMAGE_PATH))
    except Exception:
        return

    image.width = 760
    image.height = 58
    sheet.add_image(image, "A1")


def write_key_value_block(sheet, start_row: int, rows: list[tuple[str, Any]]) -> None:
    for offset, (label, value) in enumerate(rows):
        row_index = start_row + offset
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        value_cell = sheet.cell(row=row_index, column=2, value=value)
        label_cell.font = Font(bold=True, color=ACCENT_BLUE)
        label_cell.fill = PatternFill("solid", fgColor=LIGHT_FILL)
        label_cell.border = TABLE_BORDER
        value_cell.border = TABLE_BORDER
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)


def style_section_title(cell) -> None:
    cell.font = Font(bold=True, size=13, color=ACCENT_RED)
    cell.alignment = Alignment(vertical="center")


def write_table(
    sheet,
    start_row: int,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT_GREEN)
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_offset, row_values in enumerate(rows, start=1):
        row_index = start_row + row_offset
        for column_index, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def autosize_sheet(sheet) -> None:
    max_lengths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            value_length = len(str(cell.value))
            max_lengths[cell.column] = max(max_lengths.get(cell.column, 0), value_length)

    for column_index, value_length in max_lengths.items():
        width = min(max(value_length + 2, 12), 42)
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def generate_word_report(
    analytics: dict[str, Any],
    grade_config: list[dict[str, Any]],
    threshold: float,
) -> tuple[bytes, str]:
    context = build_report_context(analytics, grade_config, threshold)
    try:
        file_bytes = generate_word_report_python(context)
    except (ImportError, ModuleNotFoundError):
        try:
            file_bytes = generate_word_report_bundled_python(context)
        except RuntimeError:
            file_bytes = generate_word_report_node(context)
    return file_bytes, word_filename(context["batch"]["name"])


def generate_word_report_python(context: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    if HEADER_IMAGE_PATH.exists():
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(HEADER_IMAGE_PATH), width=Inches(6.8))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("KTU Result Analysis Report")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string(ACCENT_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(context["batch"]["name"]).bold = True
    subtitle.add_run(f"\n{context['batch']['exam_name']}")
    if context["batch"]["institution"]:
        subtitle.add_run(f"\n{context['batch']['institution']}")

    document.add_paragraph(
        f"GPA threshold used in this report: greater than {context['threshold']:.2f}"
    )

    def shade_cell(cell, fill: str) -> None:
        cell_properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        cell_properties.append(shading)

    def style_row(cells, header: bool) -> None:
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if header:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            shade_cell(cell, ACCENT_GREEN if header else "FFFFFF")

    def set_table_widths(table, widths: list[float]) -> None:
        table.autofit = False
        for column_index, width in enumerate(widths):
            table.columns[column_index].width = Inches(width)
        for row in table.rows:
            for column_index, width in enumerate(widths):
                row.cells[column_index].width = Inches(width)

    document.add_heading("Summary", level=1)
    summary_table = document.add_table(rows=0, cols=2)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.style = "Table Grid"
    summary_rows = [
        ("Batch Name", context["batch"]["name"]),
        ("Exam Name", context["batch"]["exam_name"]),
        ("Institution", context["batch"]["institution"] or "Institution not detected"),
        ("Total Students", str(context["summary"]["total_students"])),
        ("Total Departments", str(context["summary"]["total_departments"])),
        ("Total Courses", str(context["summary"]["total_courses"])),
        ("Average GPA", f"{context['summary']['average_gpa']:.2f}"),
        ("Highest GPA", f"{context['summary']['highest_gpa']:.2f}"),
        ("Lowest GPA", f"{context['summary']['lowest_gpa']:.2f}"),
        ("Students with GPA > threshold", str(len(context["threshold_students"]))),
    ]
    for label, value in summary_rows:
        cells = summary_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        style_row(cells, header=False)
    set_table_widths(summary_table, [2.1, 4.25])

    document.add_page_break()
    document.add_heading("Subject-wise Pass Percentage", level=1)
    subjects_by_department: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for subject in context["subjects"]:
        subjects_by_department[subject["department_name"]].append(subject)
    for department_name in sorted(subjects_by_department):
        document.add_heading(department_name, level=2)
        subject_table = document.add_table(rows=1, cols=4)
        subject_table.style = "Table Grid"
        headers = [
            "Course Code",
            "Course Name",
            "Passed / Appeared",
            "Pass %",
        ]
        for index, header in enumerate(headers):
            subject_table.rows[0].cells[index].text = header
        style_row(subject_table.rows[0].cells, header=True)
        for subject in subjects_by_department[department_name]:
            cells = subject_table.add_row().cells
            values = [
                subject["course_code"],
                subject["course_name"],
                f"{subject['passed_count']} / {subject['appeared_count']}",
                f"{subject['pass_percentage']:.2f}",
            ]
            for index, value in enumerate(values):
                cells[index].text = value
            style_row(cells, header=False)
        set_table_widths(subject_table, [1.0, 4.05, 1.0, 0.8])

    document.add_page_break()
    document.add_heading("Top 10 Ranks in Each Department", level=1)
    rankings_by_department: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in context["department_rankings"]:
        rankings_by_department[row["department_name"]].append(row)
    for department_name in sorted(rankings_by_department):
        document.add_heading(department_name, level=2)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for index, header in enumerate(["Rank", "Register No", "GPA", "Passed", "Failed"]):
            table.rows[0].cells[index].text = header
        style_row(table.rows[0].cells, header=True)
        for row in rankings_by_department[department_name]:
            cells = table.add_row().cells
            values = [
                str(row["rank"]),
                row["register_no"],
                f"{row['gpa']:.2f}",
                str(row["passed_count"]),
                str(row["failed_count"]),
            ]
            for index, value in enumerate(values):
                cells[index].text = value
            style_row(cells, header=False)
        set_table_widths(table, [0.65, 1.55, 0.8, 0.75, 0.75])

    document.add_page_break()
    document.add_heading("Students Who Failed", level=1)
    failed_rows = context["failed_students"] or [
        {
            "department_name": "No failed students found",
            "register_no": "",
            "gpa": "",
            "failed_subjects_text": "",
        }
    ]
    failed_by_department: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in failed_rows:
        failed_by_department[row["department_name"]].append(row)
    for department_name in sorted(failed_by_department):
        document.add_heading(department_name, level=2)
        for row in failed_by_department[department_name]:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            line = row["register_no"]
            if row["gpa"] != "":
                line += f" | GPA {row['gpa']:.2f}"
            if row["failed_subjects_text"]:
                line += f" | {row['failed_subjects_text']}"
            paragraph.add_run(line)

    document.add_page_break()
    document.add_heading(
        f"Students with GPA Greater Than {context['threshold']:.2f}",
        level=1,
    )
    document.add_paragraph(
        f"Overall count: {len(context['threshold_students'])} students"
    )

    count_table = document.add_table(rows=1, cols=2)
    count_table.style = "Table Grid"
    count_table.rows[0].cells[0].text = "Department"
    count_table.rows[0].cells[1].text = "Count"
    style_row(count_table.rows[0].cells, header=True)
    count_rows = context["threshold_counts"] or [
        {"department_name": "No students exceed the threshold", "count": 0}
    ]
    for row in count_rows:
        cells = count_table.add_row().cells
        cells[0].text = row["department_name"]
        cells[1].text = str(row["count"])
        style_row(cells, header=False)
    set_table_widths(count_table, [4.8, 0.9])

    document.add_paragraph()
    detail_table = document.add_table(rows=1, cols=5)
    detail_table.style = "Table Grid"
    for index, header in enumerate(["Register No", "Department", "GPA", "Passed", "Failed"]):
        detail_table.rows[0].cells[index].text = header
    style_row(detail_table.rows[0].cells, header=True)

    detail_rows = context["threshold_students"] or [
        {
            "register_no": "No students exceed the threshold",
            "department_name": "",
            "gpa": "",
            "passed_count": "",
            "failed_count": "",
        }
    ]
    for row in detail_rows:
        cells = detail_table.add_row().cells
        values = [
            row["register_no"],
            row["department_name"],
            "" if row["gpa"] == "" else f"{row['gpa']:.2f}",
            str(row["passed_count"]),
            str(row["failed_count"]),
        ]
        for index, value in enumerate(values):
            cells[index].text = value
        style_row(cells, header=False)
    set_table_widths(detail_table, [1.45, 2.6, 0.7, 0.75, 0.75])

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def generate_word_report_bundled_python(context: dict[str, Any]) -> bytes:
    python_binary = resolve_bundled_python_binary()
    if not python_binary:
        raise RuntimeError("Bundled Python runtime is not available.")
    if not WORD_BUILDER_PY_PATH.exists():
        raise RuntimeError("Python Word report builder script is missing.")

    with tempfile.TemporaryDirectory(dir=BASE_DIR) as temp_dir:
        temp_path = Path(temp_dir)
        context_path = temp_path / "report-context.json"
        output_path = temp_path / "analysis-report.docx"
        context_path.write_text(json.dumps(context), encoding="utf-8")

        command = [
            python_binary,
            str(WORD_BUILDER_PY_PATH),
            str(context_path),
            str(output_path),
            str(HEADER_IMAGE_PATH),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
            cwd=BASE_DIR,
        )
        if completed.returncode != 0:
            error_output = completed.stderr.strip() or completed.stdout.strip()
            raise RuntimeError(
                f"Unable to generate the Word report with bundled Python: {error_output or 'unknown error'}"
            )

        return output_path.read_bytes()


def resolve_node_binary() -> str:
    env_path = os.environ.get("RESULT_ANALYSER_NODE_BINARY")
    if env_path:
        return env_path

    return "node"


def generate_word_report_node(context: dict[str, Any]) -> bytes:
    if not WORD_BUILDER_PATH.exists():
        raise RuntimeError("Word report builder script is missing.")

    with tempfile.TemporaryDirectory(dir=BASE_DIR) as temp_dir:
        temp_path = Path(temp_dir)
        context_path = temp_path / "report-context.json"
        output_path = temp_path / "analysis-report.docx"
        context_path.write_text(json.dumps(context), encoding="utf-8")

        command = [
            resolve_node_binary(),
            str(WORD_BUILDER_PATH),
            str(context_path),
            str(output_path),
            str(HEADER_IMAGE_PATH),
        ]
        env = os.environ.copy()
        node_modules_dir = resolve_node_modules_dir()
        if node_modules_dir:
            env["RESULT_ANALYSER_NODE_MODULES"] = node_modules_dir
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
            cwd=BASE_DIR,
            env=env,
        )
        if completed.returncode != 0:
            error_output = completed.stderr.strip() or completed.stdout.strip()
            raise RuntimeError(
                f"Unable to generate the Word report: {error_output or 'unknown error'}"
            )

        return normalise_word_media_extensions(output_path.read_bytes())
